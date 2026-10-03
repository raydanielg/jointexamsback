"""Format renderers for normalized report sections."""
import csv
import html
import io

CONTENT_TYPES = {
    "PDF": "application/pdf",
    "XLSX": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    "CSV": "text/csv",
    "HTML": "text/html; charset=utf-8",
}
EXTENSIONS = {"PDF": "pdf", "XLSX": "xlsx", "CSV": "csv", "HTML": "html"}


def render(doc, fmt):
    if fmt == "PDF":
        return _render_pdf(doc)
    if fmt == "XLSX":
        return _render_xlsx(doc)
    if fmt == "CSV":
        return _render_csv(doc)
    return _render_html(doc)


def _render_pdf(doc):
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4, landscape
    from reportlab.lib.styles import getSampleStyleSheet
    from reportlab.lib.units import mm
    from reportlab.platypus import (
        PageBreak,
        Paragraph,
        SimpleDocTemplate,
        Spacer,
        Table,
        TableStyle,
    )

    buffer = io.BytesIO()
    pagesize = landscape(A4) if any(
        len(s.get("columns", [])) > 8 for s in doc["sections"]
    ) else A4
    pdf = SimpleDocTemplate(
        buffer, pagesize=pagesize, leftMargin=15 * mm, rightMargin=15 * mm,
        topMargin=15 * mm, bottomMargin=15 * mm,
    )
    styles = getSampleStyleSheet()
    story = [
        Paragraph(doc["exam_name"], styles["Title"]),
        Paragraph(f"{doc['exam_code']} — {doc['title']}", styles["Heading2"]),
        Spacer(1, 6 * mm),
    ]
    for section in doc["sections"]:
        story.append(Paragraph(section["title"], styles["Heading3"]))
        meta = section.get("meta") or {}
        if meta:
            meta_text = "  |  ".join(
                f"{k.replace('_', ' ').title()}: {html.escape(str(v))}"
                for k, v in meta.items() if v not in (None, "")
            )
            story.append(Paragraph(meta_text, styles["Normal"]))
            story.append(Spacer(1, 2 * mm))
        table_data = [section["columns"]] + [[str(c) for c in row] for row in section["rows"]]
        table = Table(table_data, repeatRows=1)
        table.setStyle(
            TableStyle(
                [
                    ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#1f3864")),
                    ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
                    ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
                    ("FONTSIZE", (0, 0), (-1, -1), 8),
                    ("GRID", (0, 0), (-1, -1), 0.25, colors.grey),
                    ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#f2f5fa")]),
                    ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
                ]
            )
        )
        story.append(table)
        story.append(Spacer(1, 8 * mm))
    pdf.build(story)
    return buffer.getvalue()


def _render_xlsx(doc):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill

    wb = Workbook()
    wb.remove(wb.active)
    header_fill = PatternFill("solid", start_color="1F3864")
    header_font = Font(bold=True, color="FFFFFF")
    for i, section in enumerate(doc["sections"]):
        title = "".join(c if c.isalnum() or c in " -" else "" for c in section["title"])[:31] or f"Section{i+1}"
        ws = wb.create_sheet(title=title)
        ws.append([doc["exam_name"], f"{doc['exam_code']} — {doc['title']}"])
        ws.append([])
        for k, v in (section.get("meta") or {}).items():
            if v not in (None, ""):
                ws.append([k.replace("_", " ").title(), str(v)])
        ws.append(section["columns"])
        for cell in ws[ws.max_row]:
            cell.font = header_font
            cell.fill = header_fill
        for row in section["rows"]:
            ws.append(["" if c is None else c for c in row])
    if not doc["sections"]:
        wb.create_sheet("Empty").append(["No data"])
    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


def _render_csv(doc):
    buffer = io.StringIO()
    writer = csv.writer(buffer)
    for section in doc["sections"]:
        writer.writerow([section["title"]])
        writer.writerow(section["columns"])
        writer.writerows(section["rows"])
        writer.writerow([])
    return buffer.getvalue().encode("utf-8")


def _render_html(doc):
    parts = [
        "<!doctype html><html><head><meta charset='utf-8'>",
        f"<title>{html.escape(doc['title'])}</title>",
        "<style>body{font-family:system-ui,sans-serif;margin:2rem;color:#111}"
        "table{border-collapse:collapse;width:100%;margin:1rem 0;font-size:.9rem}"
        "th{background:#1f3864;color:#fff;text-align:left}"
        "th,td{border:1px solid #bbb;padding:.35rem .5rem}"
        "h2{margin-bottom:0}@media print{.no-print{display:none}}</style></head><body>",
        f"<h2>{html.escape(doc['exam_name'])}</h2>",
        f"<p><strong>{html.escape(doc['exam_code'])}</strong> — {html.escape(doc['title'])}</p>",
    ]
    for section in doc["sections"]:
        parts.append(f"<h3>{html.escape(section['title'])}</h3>")
        meta = section.get("meta") or {}
        if meta:
            parts.append("<p>" + " &middot; ".join(
                f"{html.escape(k.replace('_',' ').title())}: <strong>{html.escape(str(v))}</strong>"
                for k, v in meta.items() if v not in (None, "")
            ) + "</p>")
        parts.append("<table><thead><tr>")
        parts.extend(f"<th>{html.escape(str(c))}</th>" for c in section["columns"])
        parts.append("</tr></thead><tbody>")
        for row in section["rows"]:
            parts.append("<tr>" + "".join(f"<td>{html.escape(str(c))}</td>" for c in row) + "</tr>")
        parts.append("</tbody></table>")
    parts.append("</body></html>")
    return "".join(parts).encode("utf-8")
