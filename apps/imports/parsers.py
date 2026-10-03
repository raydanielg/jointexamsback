"""Row-level spreadsheet parsing: CSV + XLSX -> list of dicts keyed by
normalized header names."""
import csv
import io

import openpyxl
from openpyxl import Workbook

from apps.core.exceptions import BusinessRuleError


def parse_spreadsheet(file_obj, filename):
    ext = filename.rsplit(".", 1)[-1].lower() if "." in filename else ""
    if ext == "xlsx":
        return _parse_xlsx(file_obj)
    if ext == "csv":
        return _parse_csv(file_obj)
    raise BusinessRuleError("Unsupported file type. Upload CSV or XLSX.", code="INVALID_FILE_TYPE")


def _norm(header):
    return str(header or "").strip().lower().replace(" ", "_").replace("-", "_")


def _parse_csv(file_obj):
    raw = file_obj.read()
    text = raw.decode("utf-8-sig", errors="replace")
    reader = csv.DictReader(io.StringIO(text))
    reader.fieldnames = [_norm(h) for h in (reader.fieldnames or [])]
    return [{k: (v.strip() if isinstance(v, str) else v) for k, v in row.items()} for row in reader]


def _parse_xlsx(file_obj):
    wb = openpyxl.load_workbook(file_obj, read_only=True, data_only=True)
    ws = wb.active
    rows = list(ws.iter_rows(values_only=True))
    if not rows:
        return []
    headers = [_norm(h) for h in rows[0]]
    data = []
    for raw_row in rows[1:]:
        if all(cell is None or str(cell).strip() == "" for cell in raw_row):
            continue
        data.append(
            {
                headers[i]: (str(raw_row[i]).strip() if raw_row[i] is not None else "")
                for i in range(len(headers))
                if headers[i]
            }
        )
    return data


def write_error_report(errors):
    """Return XLSX bytes for a list of {row, field, message} errors."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Import Errors"
    ws.append(["Row", "Field", "Error"])
    for err in errors:
        ws.append([err.get("row"), err.get("field", ""), err.get("message", "")])
    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


def build_template(columns, example=None):
    """Generate an XLSX import template."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Template"
    ws.append(columns)
    if example:
        ws.append(example)
    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()
